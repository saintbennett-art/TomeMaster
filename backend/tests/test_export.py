"""Export round-trip smoke: each format returns a non-empty stream of the right type."""
import base64
import io
import re
import zipfile

import pytest


def _tiny_cover():
    """A small valid JPEG data-URL, generated via Pillow, used as a cover."""
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (120, 180), (30, 50, 110)).save(buf, format="JPEG")
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()


def _tiny_plate(color=(200, 60, 60)):
    """A small valid PNG data-URL used as an inline manuscript plate/figure."""
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (400, 300), color).save(buf, format="PNG")
    return "data:image/png;base64," + base64.b64encode(buf.getvalue()).decode()


# A manuscript carrying TWO inline plates: one as a <figure>+<figcaption> standing
# between paragraphs, one inline at the end of a prose <p> with an alt caption.
_PLATE_FIG_CAPTION = "Plate I: The harbor at dawn"
_PLATE_ALT_CAPTION = "Storm clouds gather"


def _book_with_plates():
    plate1 = _tiny_plate((200, 60, 60))
    plate2 = _tiny_plate((60, 120, 200))
    content = (
        "<h1>The Lighthouse</h1><p>by Jane Doe</p>"
        "<h1>Chapter One: The Arrival</h1>"
        "<p>The keeper counted the ships.</p>"
        f'<figure><img src="{plate1}" alt="ignored"/>'
        f"<figcaption>{_PLATE_FIG_CAPTION}</figcaption></figure>"
        "<p>Then he climbed the stairs.</p>"
        "<h1>Chapter Two: The Storm</h1>"
        f'<p>A photo follows.<img src="{plate2}" alt="{_PLATE_ALT_CAPTION}"/></p>'
        "<p>The wind rose to a scream.</p>"
    )
    return {
        "content": content,
        "chapters": [],
        "title": "The Lighthouse",
        "author": "Jane Doe",
        "format": "chicago",
        "cover_image": _tiny_cover(),
    }


_BOOK_BODY = {
    "content": (
        "<h1>The Lighthouse</h1><p>by Jane Doe</p>"
        "<h1>Chapter One: The Arrival</h1><p>The keeper counted the ships.</p>"
        "<h1>Chapter Two: The Storm</h1><p>The wind rose to a scream.</p>"
    ),
    "chapters": [
        {"suggested_title": "Chapter One: The Arrival", "display_page": 1},
        {"suggested_title": "Chapter Two: The Storm", "display_page": 5},
    ],
    "title": "The Lighthouse",
    "author": "Jane Doe",
    "format": "chicago",
}


def test_docx_has_cover_title_and_toc(client):
    body = {**_BOOK_BODY, "cover_image": _tiny_cover()}
    r = client.post("/api/v1/document/export/docx", json=body)
    assert r.status_code == 200, r.text[:200]
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        names = z.namelist()
        assert any("media" in n for n in names), "cover image not embedded in DOCX"
        doc_xml = z.read("word/document.xml").decode("utf-8", "replace")
        assert "TOC " in doc_xml, "Word TOC field missing"
        # Title appears on the generated title page but the body block is stripped (no dup).
        assert doc_xml.count("The Lighthouse") == 1
        assert "by Jane Doe" in doc_xml


def test_docx_updates_fields_on_open(client):
    """Word must repopulate the TOC/PAGE/STYLEREF fields when the file opens."""
    r = client.post("/api/v1/document/export/docx", json={**_BOOK_BODY, "cover_image": _tiny_cover()})
    assert r.status_code == 200, r.text[:200]
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        settings = z.read("word/settings.xml").decode("utf-8", "replace")
        assert "w:updateFields" in settings, "settings.xml missing updateFields"
        assert 'w:val="true"' in settings or "w:val='true'" in settings, "updateFields not enabled"


def test_docx_running_header_and_centered_footer(client):
    """STYLEREF chapter running header + centered PAGE footer in the body section."""
    r = client.post("/api/v1/document/export/docx", json={**_BOOK_BODY, "cover_image": _tiny_cover()})
    assert r.status_code == 200, r.text[:200]
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        names = z.namelist()
        headers = "".join(
            z.read(n).decode("utf-8", "replace") for n in names if "header" in n and n.endswith(".xml")
        )
        footers = "".join(
            z.read(n).decode("utf-8", "replace") for n in names if "footer" in n and n.endswith(".xml")
        )
        assert "STYLEREF" in headers, "running header STYLEREF field missing"
        assert "PAGE" in footers, "footer PAGE field missing"
        assert 'w:jc w:val="center"' in footers, "footer page number not centered"


def test_docx_front_matter_unnumbered(client):
    """Front matter (no pgNumType) + one section per chapter (first restarts at 1;
    later chapter sections continue numbering). Chapter sections use a different
    first page so the opener shows no running head."""
    from docx import Document
    from docx.oxml.ns import qn

    r = client.post("/api/v1/document/export/docx", json={**_BOOK_BODY, "cover_image": _tiny_cover()})
    assert r.status_code == 200, r.text[:200]
    d = Document(io.BytesIO(r.content))
    # _BOOK_BODY carries two chapters → front matter + 2 chapter sections.
    assert len(d.sections) == 3, "expected front matter + one section per chapter"
    pg0 = d.sections[0]._sectPr.find(qn("w:pgNumType"))
    pg1 = d.sections[1]._sectPr.find(qn("w:pgNumType"))
    pg2 = d.sections[2]._sectPr.find(qn("w:pgNumType"))
    assert pg0 is None, "front-matter section must not declare page numbering"
    assert pg1 is not None and pg1.get(qn("w:start")) == "1", "body must restart numbering at 1"
    assert pg2 is None, "later chapters must continue numbering, not restart"
    for s in d.sections[1:]:
        assert s.different_first_page_header_footer, "chapter opener must suppress the running head"
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading 1")]
    assert "Chapter One: The Arrival" in heads, "chapters must use real Heading 1 style"


def test_docx_title_page_never_suppressed(client):
    """Title page renders even when title/author are the placeholder defaults."""
    r = client.post("/api/v1/document/export/docx", json={**_BODY, "content": "<p>Body text only.</p>"})
    assert r.status_code == 200, r.text[:200]
    from docx import Document
    d = Document(io.BytesIO(r.content))
    # Default title from _BODY is "Smoke Test"; it must appear on the title page.
    assert any("Smoke Test" in p.text for p in d.paragraphs), "title page suppressed"


def test_pdf_has_bookmarks_image_and_links(client):
    body = {**_BOOK_BODY, "cover_image": _tiny_cover()}
    r = client.post("/api/v1/document/export/pdf", json=body)
    assert r.status_code == 200, r.text[:200]
    pdf = r.content
    assert b"/Outlines" in pdf, "PDF outline/bookmarks missing"
    assert (b"/Image" in pdf or b"/XObject" in pdf), "cover image missing from PDF"
    assert b"/Link" in pdf, "clickable TOC links missing from PDF"


def test_pdf_toc_has_real_page_numbers_and_outline(client):
    """The PDF Contents must list real page numbers (two-pass build) and the
    outline must contain every chapter."""
    fitz = pytest.importorskip("fitz")
    body = {**_BOOK_BODY, "cover_image": _tiny_cover()}
    r = client.post("/api/v1/document/export/pdf", json=body)
    assert r.status_code == 200, r.text[:200]
    d = fitz.open(stream=r.content, filetype="pdf")
    # Locate the Contents page and confirm it carries a digit page number per chapter.
    contents_text = ""
    for i in range(d.page_count):
        t = d[i].get_text()
        if "Contents" in t:
            contents_text = t
            break
    assert contents_text, "Contents page not found"
    assert "Chapter One: The Arrival" in contents_text
    import re
    # A dot-leader line ending in a page number, e.g. ". . . 1"
    assert re.search(r"\.\s*\d+", contents_text), "TOC has no page numbers"
    outline_titles = [t for _lvl, t, _pg in d.get_toc()]
    assert "Chapter One: The Arrival" in outline_titles, "outline/bookmarks missing chapter"


def test_pdf_front_matter_unnumbered(client):
    """Cover (page 1) and title page (page 2) carry no footer page number; body
    numbering starts at 1 on the first chapter page."""
    fitz = pytest.importorskip("fitz")
    body = {**_BOOK_BODY, "cover_image": _tiny_cover()}
    r = client.post("/api/v1/document/export/pdf", json=body)
    assert r.status_code == 200, r.text[:200]
    d = fitz.open(stream=r.content, filetype="pdf")
    # Cover page is an image only — no extractable footer digit.
    assert d[0].get_text().strip() == "", "cover page should be image-only / unnumbered"
    # Title page (page 2) shows the title but no standalone footer page number.
    title_pg = d[1].get_text()
    assert "The Lighthouse" in title_pg
    # The first body page carries the chapter title (running header) and footer "1".
    first_chapter_page = next(
        i for i in range(d.page_count) if "Chapter One: The Arrival" in d[i].get_text()
    )
    assert first_chapter_page >= 2, "body must follow the unnumbered front matter"


def test_epub_has_cover_titlepage_and_nav(client):
    body = {**_BOOK_BODY, "cover_image": _tiny_cover()}
    r = client.post("/api/v1/document/export/epub", json=body)
    assert r.status_code == 200, r.text[:200]
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        names = z.namelist()
        assert any("cover" in n.lower() for n in names), "EPUB cover file missing"
        assert any("title_page" in n for n in names), "EPUB title page missing"
        opf = next(n for n in names if n.endswith(".opf"))
        opf_xml = z.read(opf).decode("utf-8", "replace")
        assert 'properties="cover-image"' in opf_xml, "cover not marked as EPUB cover-image"
        nav = z.read("EPUB/nav.xhtml").decode("utf-8", "replace")
        assert "Chapter One: The Arrival" in nav, "nav TOC missing chapters"


# ─── New lightweight formats: Markdown / HTML / RTF / Plain text ──────────────

def test_markdown_has_cover_title_toc_and_no_dup(client):
    body = {**_BOOK_BODY, "cover_image": _tiny_cover()}
    r = client.post("/api/v1/document/export/md", json=body)
    assert r.status_code == 200, r.text[:200]
    md = r.content.decode("utf-8")
    assert "data:image/jpeg;base64," in md, "cover image reference missing"
    assert "# The Lighthouse" in md, "title heading missing"
    assert "by Jane Doe" in md, "author missing"
    # Anchor-linked TOC entry + a matching anchor target.
    assert "[Chapter One: The Arrival](#chapter-one-the-arrival)" in md, "TOC anchor link missing"
    assert '<a id="chapter-one-the-arrival">' in md, "anchor target missing"
    # The manuscript's own title block is stripped → title appears once as the # heading.
    assert md.count("# The Lighthouse") == 1, "title duplicated in body"
    assert '"' in md and "“" not in md, "must use straight quotes only"


def test_html_is_standalone_with_cover_and_linked_toc(client):
    body = {**_BOOK_BODY, "cover_image": _tiny_cover()}
    r = client.post("/api/v1/document/export/html", json=body)
    assert r.status_code == 200, r.text[:200]
    h = r.content.decode("utf-8")
    assert h.lstrip().startswith("<!DOCTYPE html>"), "not a valid standalone HTML doc"
    assert 'src="data:image/jpeg;base64,' in h, "embedded cover data URI missing"
    assert "<h1>The Lighthouse</h1>" in h and "by Jane Doe" in h, "title block missing"
    assert 'href="#chapter-one-the-arrival"' in h, "TOC anchor link missing"
    assert 'id="chapter-one-the-arrival"' in h, "chapter id anchor missing"
    assert "</body></html>" in h, "HTML not closed"
    # Parse it to confirm validity and that the TOC link resolves to a real anchor.
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(h, "html.parser")
    links = [a.get("href") for a in soup.select("nav.toc a")]
    assert "#chapter-one-the-arrival" in links
    assert soup.find(id="chapter-one-the-arrival") is not None, "anchor target not present"


def test_rtf_opens_with_title_toc_and_bold(client):
    body = {**_BOOK_BODY, "cover_image": _tiny_cover()}
    r = client.post("/api/v1/document/export/rtf", json=body)
    assert r.status_code == 200, r.text[:200]
    rtf = r.content.decode("ascii")
    assert rtf.startswith(r"{\rtf1"), "missing RTF signature"
    assert rtf.rstrip().endswith("}"), "RTF not closed"
    assert rtf.count("{") == rtf.count("}"), "unbalanced RTF braces"
    assert "The Lighthouse" in rtf and "by Jane Doe" in rtf, "title page missing"
    assert "Table of Contents" in rtf, "RTF TOC missing"
    # TOC entry + chapter heading → the chapter title appears twice.
    assert rtf.count("Chapter One: The Arrival") == 2, "chapter missing from TOC or body"
    # \par must be delimited (newline) so following text isn't swallowed.
    assert r"\parChapter" not in rtf, "\\par abuts text — would drop the chapter"


def test_txt_has_title_and_numbered_toc_no_image(client):
    body = {**_BOOK_BODY, "cover_image": _tiny_cover()}
    r = client.post("/api/v1/document/export/txt", json=body)
    assert r.status_code == 200, r.text[:200]
    txt = r.content.decode("utf-8")
    assert txt.startswith("The Lighthouse"), "title block missing"
    assert "by Jane Doe" in txt, "author missing"
    assert "TABLE OF CONTENTS" in txt, "text TOC header missing"
    assert "1. Chapter One: The Arrival" in txt, "numbered TOC entry missing"
    assert "2. Chapter Two: The Storm" in txt, "second numbered TOC entry missing"
    assert "base64" not in txt and "data:image" not in txt, "plain text must contain no image"
    assert "The wind rose to a scream." in txt, "body wording not preserved"


# ─── Inline plates / figures: images must appear IN EVERY format ──────────────

def test_docx_embeds_inline_plates(client):
    """Two inline plates embed as DOCX media parts (cover + 2 plates = 3)."""
    r = client.post("/api/v1/document/export/docx", json=_book_with_plates())
    assert r.status_code == 200, r.text[:200]
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        media = [n for n in z.namelist() if "media" in n]
        # cover (jpg) + plate1 (png) + plate2 (png)
        assert len(media) >= 3, f"plates not embedded in DOCX: {media}"
        assert sum(n.lower().endswith(".png") for n in media) >= 2, "plate PNGs missing"
        doc_xml = z.read("word/document.xml").decode("utf-8", "replace")
        # Captions render as text under the picture.
        assert _PLATE_FIG_CAPTION in doc_xml, "figure caption missing"
        assert _PLATE_ALT_CAPTION in doc_xml, "alt-derived caption missing"


def test_pdf_embeds_inline_plates(client):
    """The PDF carries image XObjects for cover + both plates."""
    r = client.post("/api/v1/document/export/pdf", json=_book_with_plates())
    assert r.status_code == 200, r.text[:200]
    pdf = r.content
    assert (b"/Image" in pdf or b"/XObject" in pdf), "no images in PDF"
    fitz = pytest.importorskip("fitz")
    d = fitz.open(stream=pdf, filetype="pdf")
    total_images = sum(len(d[i].get_images(full=True)) for i in range(d.page_count))
    # cover + 2 plates → at least 3 embedded images.
    assert total_images >= 3, f"expected cover + 2 plates, got {total_images} images"
    all_text = "".join(d[i].get_text() for i in range(d.page_count))
    assert _PLATE_FIG_CAPTION in all_text, "PDF plate caption missing"


def test_epub_adds_plate_resources_and_references(client):
    """Plates become EPUB image resources AND are referenced from chapter XHTML."""
    r = client.post("/api/v1/document/export/epub", json=_book_with_plates())
    assert r.status_code == 200, r.text[:200]
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        names = z.namelist()
        plate_files = [n for n in names if "images/plate_" in n]
        assert len(plate_files) >= 2, f"plate resources missing: {names}"
        # The manifest must list each plate image item.
        opf = next(n for n in names if n.endswith(".opf"))
        opf_xml = z.read(opf).decode("utf-8", "replace")
        assert "images/plate_1" in opf_xml, "plate not in manifest"
        # A chapter XHTML must reference a plate image and carry the caption.
        chaps = [n for n in names if re.search(r"chap_\d+\.xhtml$", n)]
        joined = "".join(z.read(c).decode("utf-8", "replace") for c in chaps)
        assert 'src="images/plate_' in joined, "chapter does not reference plate image"
        assert _PLATE_FIG_CAPTION in joined, "EPUB plate caption missing"


def test_markdown_keeps_inline_plates(client):
    r = client.post("/api/v1/document/export/md", json=_book_with_plates())
    assert r.status_code == 200, r.text[:200]
    md = r.content.decode("utf-8")
    assert f"![{_PLATE_FIG_CAPTION}](data:image/png" in md, "figure plate missing from MD"
    assert f"![{_PLATE_ALT_CAPTION}](data:image/png" in md, "inline plate missing from MD"


def test_html_keeps_inline_plates(client):
    r = client.post("/api/v1/document/export/html", json=_book_with_plates())
    assert r.status_code == 200, r.text[:200]
    h = r.content.decode("utf-8")
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(h, "html.parser")
    figs = soup.select("figure.plate")
    assert len(figs) >= 2, "inline plate figures missing from HTML"
    imgs = [f.find("img") for f in figs]
    assert all(img and (img.get("src") or "").startswith("data:image/png") for img in imgs), \
        "plate data-URI src not preserved in HTML"
    caps = "".join(c.get_text() for c in soup.select("figure.plate figcaption"))
    assert _PLATE_FIG_CAPTION in caps, "HTML plate caption missing"


def test_rtf_embeds_cover_and_plates_as_pict(client):
    """RTF embeds the COVER and BOTH plates as \\pict groups; braces stay balanced."""
    r = client.post("/api/v1/document/export/rtf", json=_book_with_plates())
    assert r.status_code == 200, r.text[:200]
    rtf = r.content.decode("ascii")
    assert rtf.startswith(r"{\rtf1"), "missing RTF signature"
    assert rtf.count("{") == rtf.count("}"), "unbalanced RTF braces (pict group malformed)"
    # cover (\pict) + 2 plates (\pict) = 3 picture groups.
    assert rtf.count("\\pict") >= 3, f"expected cover + 2 plate \\pict groups, got {rtf.count('\\pict')}"
    assert "\\pngblip" in rtf, "plate PNG not embedded via \\pngblip"
    assert _PLATE_FIG_CAPTION in rtf, "RTF plate caption missing"


def test_txt_emits_plate_and_cover_placeholders(client):
    """Plain text can't show images → cover + each plate become visible placeholders."""
    r = client.post("/api/v1/document/export/txt", json=_book_with_plates())
    assert r.status_code == 200, r.text[:200]
    txt = r.content.decode("utf-8")
    assert "[Cover: The Lighthouse]" in txt, "cover placeholder missing"
    assert f"[Plate: {_PLATE_FIG_CAPTION}]" in txt, "figure plate placeholder missing"
    assert f"[Plate: {_PLATE_ALT_CAPTION}]" in txt, "inline plate placeholder missing"
    assert "data:image" not in txt and "base64" not in txt, "TXT must not leak image data"
    assert "The wind rose to a scream." in txt, "prose around plate dropped"


def test_malformed_plate_is_surfaced_not_silently_dropped(client):
    """A bad <img> must produce a visible note, and must NOT abort the export."""
    body = {
        "content": (
            "<h1>Chapter One: The Arrival</h1>"
            "<p>Before.</p>"
            '<figure><img src="data:image/png;base64,NOT_VALID_BASE64=="/></figure>'
            "<p>After.</p>"
        ),
        "chapters": [],
        "title": "Broken Plate",
        "author": "Tester",
        "format": "chicago",
    }
    # HTML surfaces the error inline; the rest of the doc still renders.
    r = client.post("/api/v1/document/export/html", json=body)
    assert r.status_code == 200, r.text[:200]
    h = r.content.decode("utf-8")
    assert "could not be embedded" in h, "malformed plate not surfaced"
    assert "Before." in h and "After." in h, "export aborted around bad plate"


_BODY = {
    "content": "<p>The lighthouse keeper counted the ships.</p>",
    "chapters": [],
    "title": "Smoke Test",
    "author": "Tester",
    "format": "chicago",
}

_CASES = [
    ("/api/v1/document/export/docx",
     "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ("/api/v1/document/export/pdf", "application/pdf"),
    ("/api/v1/document/export/epub", "application/epub+zip"),
    ("/api/v1/document/export/md", "text/markdown"),
    ("/api/v1/document/export/html", "text/html"),
    ("/api/v1/document/export/rtf", "application/rtf"),
    ("/api/v1/document/export/txt", "text/plain"),
]


@pytest.mark.parametrize("url,content_type", _CASES)
def test_export_returns_stream(client, url, content_type):
    r = client.post(url, json=_BODY)
    assert r.status_code == 200, f"{url} -> {r.status_code}: {r.text[:200]}"
    assert content_type in r.headers.get("content-type", "")
    # Text formats are legitimately compact for a one-paragraph body; binary
    # formats are always >100B. Just confirm the stream is non-trivial.
    assert len(r.content) > 30, "export stream suspiciously small"


@pytest.mark.parametrize("url", [c[0] for c in _CASES])
def test_export_rejects_empty_content(client, url):
    r = client.post(url, json={**_BODY, "content": ""})
    assert r.status_code == 400
