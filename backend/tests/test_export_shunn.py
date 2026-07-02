"""Shunn Standard Manuscript Format ('submission' DOCX preset) — layout assertions,
plus proof that the chicago/penguin presets are structurally unaffected."""
import base64
import io
import re
import zipfile

import pytest


def _tiny_cover():
    """A small valid JPEG data-URL, generated via Pillow, used as a cover."""
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (120, 180), (30, 50, 110)).save(buf, format="JPEG")
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()


_SUB_BODY = {
    "content": (
        "<h1>The Lighthouse</h1><p>by Jane Doe</p>"
        "<h1>Chapter One: The Arrival</h1><p>The keeper counted the ships.</p>"
        "<p>***</p>"
        "<p>Then he <em>slowly</em> climbed the stairs.</p>"
        "<h1>Chapter Two: The Storm</h1><p>The wind rose to a scream.</p>"
    ),
    "chapters": [],
    "title": "The Lighthouse",
    "author": "Jane Doe",
    "format": "submission",
}

_BOOK_BODY = {
    "content": (
        "<h1>The Lighthouse</h1><p>by Jane Doe</p>"
        "<h1>Chapter One: The Arrival</h1><p>The keeper counted the ships.</p>"
        "<h1>Chapter Two: The Storm</h1><p>The wind rose to a scream.</p>"
    ),
    "chapters": [],
    "title": "The Lighthouse",
    "author": "Jane Doe",
    "format": "chicago",
}


def _export_docx(client, body):
    r = client.post("/api/v1/document/export/docx", json=body)
    assert r.status_code == 200, r.text[:200]
    return r.content


def _zip_parts(raw):
    """(document.xml, joined header xml, joined footer xml, namelist) of a DOCX."""
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        names = z.namelist()
        doc_xml = z.read("word/document.xml").decode("utf-8", "replace")
        headers = "".join(
            z.read(n).decode("utf-8", "replace") for n in names if "header" in n and n.endswith(".xml")
        )
        footers = "".join(
            z.read(n).decode("utf-8", "replace") for n in names if "footer" in n and n.endswith(".xml")
        )
    return doc_xml, headers, footers, names


# ─── The submission preset itself ─────────────────────────────────────────────

def test_submission_normal_style_tnr12_double_spaced_ragged_right(client):
    """TNR 12pt, exactly double-spaced, left-aligned, zero paragraph spacing."""
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    d = Document(io.BytesIO(_export_docx(client, _SUB_BODY)))
    normal = d.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size == Pt(12)
    pf = normal.paragraph_format
    assert pf.line_spacing == 2.0, "manuscript must be double-spaced"
    assert pf.alignment == WD_ALIGN_PARAGRAPH.LEFT, "must be ragged right, not justified"
    assert pf.space_before == Pt(0) and pf.space_after == Pt(0), \
        "no extra space between paragraphs"
    for s in d.sections:
        from docx.shared import Inches
        assert s.left_margin == s.right_margin == s.top_margin == s.bottom_margin == Inches(1)


def test_submission_page_one_wordcount_and_byline(client):
    """Page 1: author name, an 'about N words' line, title once, 'by Author'."""
    doc_xml, _h, _f, _names = _zip_parts(_export_docx(client, _SUB_BODY))
    assert re.search(r"about [\d,]+ words", doc_xml), "approximate word count line missing"
    assert "by Jane Doe" in doc_xml, "byline missing"
    # The manuscript's own leading title block is stripped → title appears exactly
    # once in the body xml (the centered page-1 title).
    assert doc_xml.count("The Lighthouse") == 1, "title duplicated or missing"


def test_submission_wordcount_rounds_to_nearest_500(client):
    """1,240 body words → 'about 1,000 words' (nearest 500, comma-formatted)."""
    body = {
        **_SUB_BODY,
        "content": "<h1>Chapter One</h1><p>" + " ".join(["word"] * 1238) + "</p>",
    }
    # plain text = 1238 'word' tokens + 'Chapter One' heading = 1240 words
    doc_xml, _h, _f, _names = _zip_parts(_export_docx(client, body))
    assert "about 1,000 words" in doc_xml, "word count not rounded to nearest 500"


def test_submission_running_header_lastname_title_page_field(client):
    """Pages 2+: top-right 'Doe / THE LIGHTHOUSE / <PAGE field>'; page 1 header
    suppressed via the different-first-page flag."""
    doc_xml, headers, _f, _names = _zip_parts(_export_docx(client, _SUB_BODY))
    assert "titlePg" in doc_xml, "different-first-page flag missing (header would show on page 1)"
    assert "Doe / THE LIGHTHOUSE / " in headers, "running header slug missing"
    assert "PAGE" in headers, "header PAGE field missing"
    assert 'w:jc w:val="right"' in headers, "running header not right-aligned"


def test_submission_first_line_indent_scene_break_and_chapter_pages(client):
    """Body paragraphs indent 0.5\"; '***' becomes a centered '#'; each of the two
    chapters opens with a page break."""
    doc_xml, _h, _f, _names = _zip_parts(_export_docx(client, _SUB_BODY))
    assert 'w:firstLine="720"' in doc_xml, "0.5in first-line indent missing"
    assert "<w:t>#</w:t>" in doc_xml, "scene break not rendered as centered '#'"
    assert doc_xml.count('w:type="page"') >= 2, "chapters must each start on a new page"


def test_submission_keeps_italics(client):
    """Modern Shunn: italics stay italics (no underline substitution)."""
    from docx import Document
    d = Document(io.BytesIO(_export_docx(client, _SUB_BODY)))
    italic_runs = [r.text for p in d.paragraphs for r in p.runs if r.italic]
    assert "slowly" in italic_runs, "italic run not preserved"
    assert not any(r.underline for p in d.paragraphs for r in p.runs if r.text == "slowly"), \
        "italics must not be converted to underline"


def test_submission_ignores_cover_and_has_no_toc(client):
    """cover_image is ignored (no media part) and no Word TOC field is emitted."""
    body = {**_SUB_BODY, "cover_image": _tiny_cover()}
    doc_xml, _h, _f, names = _zip_parts(_export_docx(client, body))
    assert not any("media" in n for n in names), "submission preset must not embed a cover"
    assert "TOC " not in doc_xml, "submission manuscript must not carry a Contents field"


# ─── Regression: chicago / penguin structurally unaffected ────────────────────
# (Same invariants tests/test_export.py relies on — the submission branch must
# not perturb the existing presets in any way.)

def test_chicago_structure_unaffected(client):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    raw = _export_docx(client, {**_BOOK_BODY, "cover_image": _tiny_cover()})
    doc_xml, headers, footers, names = _zip_parts(raw)
    assert any("media" in n for n in names), "chicago cover no longer embedded"
    assert "TOC " in doc_xml, "chicago Word TOC field missing"
    assert doc_xml.count("The Lighthouse") == 1
    assert "by Jane Doe" in doc_xml
    assert "STYLEREF" in headers, "chicago running header STYLEREF field missing"
    assert "PAGE" in footers, "chicago footer PAGE field missing"
    assert 'w:jc w:val="center"' in footers, "chicago footer page number not centered"

    d = Document(io.BytesIO(raw))
    # Front matter + one section per chapter (chapter openers suppress the running head).
    assert len(d.sections) == 3, "chicago must keep front matter + a section per chapter"
    pg0 = d.sections[0]._sectPr.find(qn("w:pgNumType"))
    pg1 = d.sections[1]._sectPr.find(qn("w:pgNumType"))
    assert pg0 is None, "chicago front matter must stay unnumbered"
    assert pg1 is not None and pg1.get(qn("w:start")) == "1", "chicago body must restart at 1"
    normal = d.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size == Pt(12)
    assert normal.paragraph_format.line_spacing == 2.0
    assert normal.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY, \
        "chicago body must stay justified"


def test_penguin_structure_unaffected(client):
    from docx import Document

    raw = _export_docx(client, {**_BOOK_BODY, "format": "penguin"})
    d = Document(io.BytesIO(raw))
    normal = d.styles["Normal"]
    assert normal.font.name == "Garamond", "penguin must keep Garamond"
    assert normal.paragraph_format.line_spacing == 1.5, "penguin must keep 1.5 spacing"
    assert len(d.sections) == 3, "penguin must keep front matter + a section per chapter"
